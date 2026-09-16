"""Excel extractor. openpyxl when installed; otherwise a zero-dependency
reader built on stdlib ``zipfile`` + ``xml.etree`` that parses xlsx's
shared strings and the first worksheet. CSV/TSV use the stdlib csv module.
"""

from __future__ import annotations

import csv
import io
import re
import xml.etree.ElementTree as ET
import zipfile

from features.common.errors import BackendUnavailableError

from . import ExtractedDoc

_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


class ExcelExtractor:
    name = "excel"

    def extract(self, data: bytes, filename: str) -> ExtractedDoc:
        lowered = filename.lower()
        if lowered.endswith((".csv", ".tsv")):
            return self._csv(data, filename)
        if lowered.endswith((".xlsx", ".xlsm")):
            return self._xlsx(data, filename)
        raise BackendUnavailableError(
            f"unsupported spreadsheet format {filename!r}",
            hint="use .csv/.tsv/.xlsx",
        )

    # -- csv / tsv ----------------------------------------------------------

    def _csv(self, data: bytes, filename: str) -> ExtractedDoc:
        text = data.decode("utf-8", errors="replace")
        delim = "\t" if filename.lower().endswith(".tsv") else ","
        rows = list(csv.reader(io.StringIO(text), delimiter=delim))
        header = rows[0] if rows else []
        lines = ["\t".join(header)] if header else []
        for r in rows[1:]:
            lines.append("\t".join(r))
        blocks = [{"text": "\t".join(r)} for r in rows[1:]]
        return ExtractedDoc(
            text="\n".join(lines),
            blocks=blocks,
            metadata={"kind": "spreadsheet", "format": filename.rsplit(".", 1)[-1],
                      "rows": max(0, len(rows) - 1), "columns": len(header)},
        )

    # -- xlsx ---------------------------------------------------------------

    def _xlsx(self, data: bytes, filename: str) -> ExtractedDoc:
        try:
            import openpyxl  # type: ignore  (lazy by design)

            return self._openpyxl(data, filename)
        except ImportError:
            return self._raw_xlsx(data, filename)

    @staticmethod
    def _openpyxl(data: bytes, filename: str) -> ExtractedDoc:
        import io

        import openpyxl  # type: ignore

        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        blocks = []
        lines = []
        total_rows = 0
        for ws in wb.worksheets:
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                cells = ["" if v is None else str(v) for v in row]
                if not any(cells):
                    continue
                line = "\t".join(cells)
                if i == 0:
                    lines.append(f"[{ws.title}]")
                    lines.append(line)
                else:
                    lines.append(line)
                    blocks.append({"text": line, "sheet": ws.title, "row": i})
                    total_rows += 1
        wb.close()
        return ExtractedDoc(
            text="\n".join(lines), blocks=blocks,
            metadata={"kind": "spreadsheet", "format": "xlsx", "rows": total_rows,
                      "sheets": len(wb.sheetnames)},
        )

    @staticmethod
    def _raw_xlsx(data: bytes, filename: str) -> ExtractedDoc:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            shared = _shared_strings(zf)
            sheet_name = _first_sheet_path(zf)
            rows = _parse_sheet(zf.read(sheet_name), shared)
        header = rows[0] if rows else []
        lines = ["\t".join(header)] if header else []
        blocks = []
        for i, r in enumerate(rows[1:], start=1):
            line = "\t".join(r)
            lines.append(line)
            blocks.append({"text": line, "sheet": 1, "row": i})
        return ExtractedDoc(
            text="\n".join(lines), blocks=blocks,
            metadata={"kind": "spreadsheet", "format": "xlsx",
                      "rows": max(0, len(rows) - 1), "columns": len(header),
                      "backend": "stdlib"},
        )


def _shared_strings(zf: zipfile.ZipFile) -> list[str]:
    if "xl/sharedStrings.xml" not in zf.namelist():
        return []
    root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
    out = []
    for si in root.findall(f"{_NS}si"):
        text = "".join(t.text or "" for t in si.iter(f"{_NS}t"))
        out.append(text)
    return out


def _first_sheet_path(zf: zipfile.ZipFile) -> str:
    candidates = [n for n in zf.namelist() if re.match(r"xl/worksheets/sheet\d+\.xml", n)]
    if not candidates:
        raise BackendUnavailableError("no worksheet found in xlsx")
    return sorted(candidates)[0]


def _col_index(ref: str) -> int:
    m = re.match(r"([A-Z]+)", ref or "")
    if not m:
        return 0
    letters = m.group(1)
    n = 0
    for ch in letters:
        n = n * 26 + (ord(ch) - 64)
    return n - 1


def _parse_sheet(xml: bytes, shared: list[str]) -> list[list[str]]:
    root = ET.fromstring(xml)
    rows: list[list[str]] = []
    for row_el in root.iter(f"{_NS}row"):
        cells: list[tuple[int, str]] = []
        for c in row_el.iter(f"{_NS}c"):
            ref = c.get("r", "")
            idx = _col_index(ref)
            t = c.get("t")
            v_el = c.find(f"{_NS}v")
            is_el = c.find(f"{_NS}is")
            if t == "s" and v_el is not None and v_el.text is not None:
                val = shared[int(v_el.text)]
            elif t == "inlineStr" and is_el is not None:
                val = "".join(x.text or "" for x in is_el.iter(f"{_NS}t"))
            elif v_el is not None:
                val = v_el.text or ""
            else:
                continue
            cells.append((idx, val))
        if not cells:
            continue
        width = max(i for i, _ in cells) + 1
        row = [""] * width
        for i, val in cells:
            if i < width:
                row[i] = val
        rows.append(row)
    return rows
